import io

from PIL import Image as PILImage
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from character_vein_api.artifact_documents import build_rich_document
from character_vein_api.artifact_exports import (
    export_docx,
    export_html,
    export_pdf,
    export_pptx,
    export_txt,
)
from character_vein_domain import CreativeArtifactDraft, RichBlockKind


def _draft(
    content: str = "# 雷的文化故事\n\n雷声唤起人们对自然力量的想象。",
) -> CreativeArtifactDraft:
    return CreativeArtifactDraft(
        artifact_type="文化文章",
        format_label="图文文章",
        title="雷的文化故事",
        summary="从历史字形出发理解雷的文化意象。",
        content_markdown=content,
        visual_hook_ids=["glyph-vertical"],
    )


def _vertical_png() -> bytes:
    image = PILImage.new("RGB", (120, 300), "white")
    for x in range(48, 72):
        for y in range(24, 276):
            image.putpixel((x, y), (28, 24, 20))
    output = io.BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def test_rich_document_has_stable_blocks_and_all_public_formats() -> None:
    first = build_rich_document(_draft())
    revised = build_rich_document(
        _draft("# 雷的文化故事\n\n这一段经过了修改。"),
        first,
    )

    assert first.supported_formats == ["docx", "html", "txt", "pdf", "pptx"]
    assert [block.block_id for block in first.pages[0].blocks] == [
        block.block_id for block in revised.pages[0].blocks
    ]
    assert (
        sum(
            block.kind == RichBlockKind.HEADING and block.text == first.title
            for block in first.pages[0].blocks
        )
        == 1
    )
    image_block = next(
        block for block in first.pages[0].blocks if block.kind == RichBlockKind.IMAGE
    )
    assert image_block.locked is True
    assert image_block.source_kind == "glyph"


def test_long_legacy_markdown_is_split_at_third_level_headings() -> None:
    sections = "\n\n".join(
        f"### 观察 {index}\n\n第 {index} 个字形观察段落。" for index in range(1, 31)
    )
    document = build_rich_document(_draft(f"## 字形观察\n\n{sections}"))

    assert len(document.pages) == 31
    assert all(len(page.blocks) <= 3 for page in document.pages)


def test_long_legacy_markdown_without_headings_is_chunked_for_slides() -> None:
    paragraphs = "\n\n".join(f"第 {index} 个独立段落。" for index in range(1, 31))
    document = build_rich_document(_draft(paragraphs))

    assert len(document.pages) == 6
    assert all(len(page.blocks) <= 7 for page in document.pages)


def test_markdown_control_syntax_is_not_exposed_as_public_copy() -> None:
    document = build_rich_document(_draft("**开场**\n\n---\n\n- **受众**：中学生\n- 风格：启发式"))
    public_text = "\n".join(block.text for page in document.pages for block in page.blocks)

    assert "**" not in public_text
    assert "---" not in public_text
    assert "受众：中学生" in public_text


def test_five_exporters_emit_openable_files_and_preserve_image_ratio() -> None:
    document = build_rich_document(_draft())
    image = _vertical_png()
    images = {"glyph-vertical": (image, "image/png")}

    html = export_html(document, images)
    text = export_txt(document)
    docx = export_docx(document, images)
    pdf = export_pdf(document, images)
    pptx = export_pptx(document, images)

    assert html.startswith(b"<!doctype html>")
    assert b"data:image/png;base64," in html
    assert "图片说明" in text.decode("utf-8-sig")
    assert b"data:image" not in text
    assert docx.startswith(b"PK")
    assert pdf.startswith(b"%PDF")
    assert pptx.startswith(b"PK")

    deck = Presentation(io.BytesIO(pptx))
    picture = next(
        shape for shape in deck.slides[0].shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    )
    assert picture.height > picture.width * 2
