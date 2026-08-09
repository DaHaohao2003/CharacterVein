import re
from io import BytesIO

from docx import Document
from pypdf import PdfReader

from character_vein_domain import AssetAnalysis


def extract_document_text(content: bytes, mime_type: str, filename: str) -> str:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if mime_type == "application/pdf" or suffix == "pdf":
        reader = PdfReader(BytesIO(content))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif (
        mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or suffix == "docx"
    ):
        document = Document(BytesIO(content))
        blocks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            blocks.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
        text = "\n".join(blocks)
    else:
        text = _decode_text(content)
    return text.strip()[:50_000]


def analyze_document_text(text: str) -> AssetAnalysis:
    targets: list[str] = []
    for pattern in (
        r"[“\"『「]([\u4e00-\u9fff])[”\"』」]\s*字?",
        r"(?:目标字|汉字|讲解对象|研究对象)\s*[：:]?\s*[“\"『「]?([\u4e00-\u9fff])[”\"』」]?",
        r"([\u4e00-\u9fff])字(?:讲解|演化|文化|课程|故事|研究|介绍|内容)",
    ):
        targets.extend(re.findall(pattern, text))
    periods = re.findall(
        r"(?:商代|西周|东周|春秋|战国|秦代|汉代|唐代|宋代|明代|清代)(?:早期|中期|晚期)?",
        text,
    )
    source_lines = [
        line.strip()
        for line in text.splitlines()
        if re.match(r"\s*(?:来源|参考|出处|引用)\s*[：:]", line)
    ]
    theme_lines = [
        line.strip()
        for line in text.splitlines()
        if re.match(r"\s*(?:主题|选题|内容方向)\s*[：:]", line)
    ]
    return AssetAnalysis(
        target_characters=list(dict.fromkeys(targets)),
        extracted_text=text,
        explicit_periods=list(dict.fromkeys(periods)),
        themes=theme_lines[:10],
        user_sources=source_lines[:20],
        requires_target_confirmation=not targets,
        model="local-document-parser",
    )


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")
