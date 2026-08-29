from __future__ import annotations

import base64
import html
import io
from collections.abc import Mapping
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image as PILImage
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches
from pptx.util import Pt as PptPt
from reportlab.lib import colors  # type: ignore[import-untyped]
from reportlab.lib.enums import TA_CENTER  # type: ignore[import-untyped]
from reportlab.lib.pagesizes import A4  # type: ignore[import-untyped]
from reportlab.lib.styles import (  # type: ignore[import-untyped]
    ParagraphStyle,
    getSampleStyleSheet,
)
from reportlab.lib.units import cm  # type: ignore[import-untyped]
from reportlab.pdfbase import pdfmetrics  # type: ignore[import-untyped]
from reportlab.pdfbase.cidfonts import UnicodeCIDFont  # type: ignore[import-untyped]
from reportlab.pdfbase.ttfonts import TTFError, TTFont  # type: ignore[import-untyped]
from reportlab.platypus import (  # type: ignore[import-untyped]
    Image,
    PageBreakIfNotEmpty,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
)

from character_vein_domain import RichArtifactDocument, RichBlockKind


def export_txt(document: RichArtifactDocument) -> bytes:
    lines = [document.title, "=" * len(document.title), ""]
    for index, page in enumerate(document.pages, start=1):
        lines.extend([f"第 {index} 页：{page.title}", ""])
        for block in page.blocks:
            if block.kind == RichBlockKind.IMAGE:
                lines.append(f"[图片说明：{block.caption or block.alt_text or '文化配图'}]")
            elif block.text:
                lines.append(block.text)
                lines.append("")
    return "\n".join(lines).encode("utf-8-sig")


def export_html(
    document: RichArtifactDocument,
    images: Mapping[str, tuple[bytes, str]],
) -> bytes:
    pages: list[str] = []
    for page in document.pages:
        blocks: list[str] = []
        for block in page.blocks:
            image_id = block.image_id
            if block.kind == RichBlockKind.IMAGE and image_id is not None and image_id in images:
                payload, mime_type = images[image_id]
                data_url = f"data:{mime_type};base64,{base64.b64encode(payload).decode('ascii')}"
                blocks.append(
                    "<figure>"
                    f'<img src="{data_url}" alt="{html.escape(block.alt_text)}">'
                    f"<figcaption>{html.escape(block.caption)}</figcaption>"
                    "</figure>"
                )
            elif block.kind == RichBlockKind.HEADING:
                level = min(3, max(1, block.level or 2))
                blocks.append(f"<h{level}>{html.escape(block.text)}</h{level}>")
            elif block.kind == RichBlockKind.LIST:
                items = "".join(
                    f"<li>{html.escape(item)}</li>" for item in block.text.splitlines() if item
                )
                blocks.append(f"<ul>{items}</ul>")
            elif block.kind == RichBlockKind.QUOTE:
                blocks.append(f"<blockquote>{html.escape(block.text)}</blockquote>")
            elif block.text:
                blocks.append(f"<p>{html.escape(block.text)}</p>")
        pages.append(f'<section class="page">{"".join(blocks)}</section>')
    content = "".join(pages)
    return f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(document.title)}</title><style>
:root{{--ink:#2f2925;--accent:#a9412c;--paper:#fffaf2}}*{{box-sizing:border-box}}
body{{margin:0;background:#eee7dc;color:var(--ink);line-height:1.8;
font-family:"Noto Serif SC","Songti SC",serif}}
.page{{width:min(900px,calc(100% - 32px));min-height:1120px;margin:32px auto;
padding:72px;background:var(--paper);box-shadow:0 12px 40px #4b2d1d18}}
h1,h2,h3{{line-height:1.3}}h1{{font-size:2.6rem}}h2{{color:var(--accent);margin-top:2rem}}
figure{{margin:2rem 0;text-align:center}}
img{{max-width:100%;max-height:430px;object-fit:contain;background:white}}
figcaption{{color:#776d66;font-size:.86rem}}
blockquote{{border-left:4px solid var(--accent);margin:1.5rem 0;
padding:.5rem 1.25rem;background:#f4e9df}}
@media(max-width:640px){{.page{{min-height:0;margin:0;width:100%;
padding:28px 20px;box-shadow:none}}}}
</style></head><body>{content}</body></html>""".encode()


def export_docx(
    document: RichArtifactDocument,
    images: Mapping[str, tuple[bytes, str]],
) -> bytes:
    output = io.BytesIO()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.3)
    normal = doc.styles["Normal"]
    normal.font.name = "Noto Serif CJK SC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(11)
    for page_index, page in enumerate(document.pages):
        if page_index:
            doc.add_page_break()  # type: ignore[no-untyped-call]
        for block in page.blocks:
            image_id = block.image_id
            if block.kind == RichBlockKind.IMAGE and image_id is not None and image_id in images:
                payload, _ = images[image_id]
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(io.BytesIO(payload), width=Cm(11.5))
                if block.caption:
                    caption = doc.add_paragraph(block.caption)
                    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    caption.runs[0].font.size = Pt(9)
            elif block.kind == RichBlockKind.HEADING:
                doc.add_heading(block.text, level=min(3, max(1, block.level or 2)))
            elif block.kind == RichBlockKind.LIST:
                for item in block.text.splitlines():
                    if item:
                        doc.add_paragraph(item, style="List Bullet")
            elif block.kind == RichBlockKind.QUOTE:
                paragraph = doc.add_paragraph(block.text)
                paragraph.paragraph_format.left_indent = Cm(0.8)
                paragraph.runs[0].italic = True
            elif block.text:
                doc.add_paragraph(block.text)
    doc.save(output)
    return output.getvalue()


def export_pptx(
    document: RichArtifactDocument,
    images: Mapping[str, tuple[bytes, str]],
) -> bytes:
    output = io.BytesIO()
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    blank_layout = deck.slide_layouts[6]
    for page_index, page in enumerate(document.pages, start=1):
        slide = deck.slides.add_slide(blank_layout)
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor(250, 247, 240)  # type: ignore[no-untyped-call]
        title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.48), Inches(8.2), Inches(0.75))
        title_frame = title_box.text_frame
        title_frame.clear()
        title_run = title_frame.paragraphs[0].add_run()
        title_run.text = page.title
        title_run.font.name = "Noto Serif CJK SC"
        title_run.font.size = PptPt(36)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(47, 41, 37)  # type: ignore[no-untyped-call]
        text_blocks = [
            block
            for block in page.blocks
            if block.kind not in {RichBlockKind.IMAGE, RichBlockKind.HEADING} and block.text
        ]
        image_blocks = [
            block
            for block in page.blocks
            if block.kind == RichBlockKind.IMAGE and block.image_id in images
        ]
        text_width = 6.4 if image_blocks else 11.7
        body_box = slide.shapes.add_textbox(
            Inches(0.78), Inches(1.55), Inches(text_width), Inches(5.1)
        )
        frame = body_box.text_frame
        frame.word_wrap = True
        frame.clear()
        for block_index, block in enumerate(text_blocks[:7]):
            paragraph = frame.paragraphs[0] if block_index == 0 else frame.add_paragraph()
            paragraph.text = (
                "• " + block.text.replace("\n", "\n• ")
                if block.kind == RichBlockKind.LIST
                else block.text
            )
            paragraph.font.name = "Noto Sans CJK SC"
            paragraph.font.size = PptPt(18 if len(text_blocks) <= 4 else 16)
            paragraph.font.color.rgb = RGBColor(72, 62, 55)  # type: ignore[no-untyped-call]
            paragraph.space_after = PptPt(12)
        if image_blocks:
            block = image_blocks[0]
            assert block.image_id is not None
            payload, _ = images[block.image_id]
            image_width, image_height = _contained_size(payload, 4.7, 4.8)
            image_left = 7.7 + (4.7 - image_width) / 2
            image_top = 1.55 + (4.8 - image_height) / 2
            slide.shapes.add_picture(
                io.BytesIO(payload),
                Inches(image_left),
                Inches(image_top),
                width=Inches(image_width),
                height=Inches(image_height),
            )
            caption = slide.shapes.add_textbox(Inches(7.7), Inches(6.45), Inches(4.7), Inches(0.35))
            caption_frame = caption.text_frame
            caption_frame.paragraphs[0].text = block.caption or block.alt_text
            caption_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            caption_frame.paragraphs[0].font.size = PptPt(10)
            caption_frame.paragraphs[0].font.color.rgb = RGBColor(  # type: ignore[no-untyped-call]
                112, 99, 90
            )
        footer = slide.shapes.add_textbox(Inches(11.8), Inches(7.0), Inches(0.7), Inches(0.25))
        footer.text_frame.paragraphs[0].text = str(page_index)
        footer.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
        footer.text_frame.paragraphs[0].font.size = PptPt(9)
    deck.save(output)
    return output.getvalue()


def export_pdf(
    document: RichArtifactDocument,
    images: Mapping[str, tuple[bytes, str]],
) -> bytes:
    output = io.BytesIO()
    font_name = _register_cjk_font()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "CvTitle",
        parent=styles["Title"],
        fontName=font_name,
        fontSize=24,
        leading=32,
        textColor=colors.HexColor("#2f2925"),
        alignment=TA_CENTER,
        spaceAfter=18,
    )
    body_style = ParagraphStyle(
        "CvBody",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=11,
        leading=18,
        textColor=colors.HexColor("#413833"),
        spaceAfter=10,
    )
    heading_style = ParagraphStyle(
        "CvHeading",
        parent=body_style,
        fontSize=17,
        leading=23,
        textColor=colors.HexColor("#a9412c"),
        spaceBefore=8,
        spaceAfter=12,
    )
    caption_style = ParagraphStyle(
        "CvCaption",
        parent=body_style,
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#70635a"),
        alignment=TA_CENTER,
    )
    story: list[object] = []
    for page_index, page in enumerate(document.pages):
        if page_index:
            story.append(PageBreakIfNotEmpty())
        story.append(Paragraph(html.escape(page.title), title_style))
        for block in page.blocks:
            image_id = block.image_id
            if block.kind == RichBlockKind.IMAGE and image_id is not None and image_id in images:
                payload, _ = images[image_id]
                image_width, image_height = _contained_size(
                    payload,
                    8.5 * cm,
                    7.5 * cm,
                )
                figure = Image(
                    io.BytesIO(payload),
                    width=image_width,
                    height=image_height,
                )
                figure.hAlign = "CENTER"
                story.append(figure)
                if block.caption:
                    story.append(Paragraph(html.escape(block.caption), caption_style))
            elif block.kind == RichBlockKind.HEADING and block.text != page.title:
                story.append(Paragraph(html.escape(block.text), heading_style))
            elif block.kind == RichBlockKind.LIST:
                for item in block.text.splitlines():
                    if item:
                        story.append(Paragraph(f"· {html.escape(item)}", body_style))
            elif block.text:
                text = "<br/>".join(html.escape(line) for line in block.text.splitlines())
                story.append(Paragraph(text, body_style))
            story.append(Spacer(1, 0.12 * cm))
    pdf = SimpleDocTemplate(
        output,
        pagesize=A4,
        leftMargin=2.1 * cm,
        rightMargin=2.1 * cm,
        topMargin=1.8 * cm,
        bottomMargin=1.8 * cm,
        title=document.title,
    )
    pdf.build(story)
    return output.getvalue()


def _contained_size(payload: bytes, max_width: float, max_height: float) -> tuple[float, float]:
    with PILImage.open(io.BytesIO(payload)) as image:
        pixel_width, pixel_height = image.size
    if pixel_width <= 0 or pixel_height <= 0:
        return max_width, max_height
    scale = min(max_width / pixel_width, max_height / pixel_height)
    return pixel_width * scale, pixel_height * scale


def _register_cjk_font() -> str:
    candidates = (
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("C:/Windows/Fonts/simsun.ttc"),
        Path("C:/Windows/Fonts/msyh.ttc"),
    )
    for path in candidates:
        if path.exists():
            name = "CharacterVeinCJK"
            if name not in pdfmetrics.getRegisteredFontNames():
                try:
                    pdfmetrics.registerFont(TTFont(name, str(path), subfontIndex=0))
                except TTFError:
                    continue
            return name
    name = "STSong-Light"
    if name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont(name))
    return name
